"""Tests for the PAB schedule workbook."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from scripts.build_pab_plan import HEADERS, build_workbook


def test_headers_and_overdue_rule(tmp_path: Path) -> None:
    path = build_workbook(tmp_path / "PAB产品计划表.xlsx")
    wb = load_workbook(path)
    ws = wb.active

    assert [ws.cell(1, col).value for col in range(1, 10)] == list(HEADERS)
    assert ws["A1"].value == "阶段"
    assert ws["B1"].value == "任务"
    assert ws["C1"].value == "任务简要"
    assert ws["D1"].value == "任务摘要"
    assert ws["E1"].value == "预估完成节点"
    assert ws["F1"].value == "实际完成节点"
    assert ws["G1"].value == "超时原因"
    assert ws["H1"].value == "任务责任人"
    assert ws["I1"].value == "偏移天数"

    tasks = {ws.cell(row, 3).value: row for row in range(2, 46) if ws.cell(row, 3).value}
    assert "PAB060·零件到位" in tasks
    row = tasks["PAB060·零件到位"]
    assert ws.cell(row, 5).value.date() == date(2026, 9, 24)
    assert ws.cell(row, 6).value is None
    assert ws.cell(row, 8).value == "王杰"
    assert "协助人：吴晨阳" in ws.cell(row, 4).value
    assert "待填" in str(ws.cell(row, 9).value)

    row_180 = tasks["PAB180·零件到位"]
    assert ws.cell(row_180, 5).value.date() == date(2026, 10, 15)
    assert "国庆假期" in ws.cell(row_180, 4).value

    formulas = []
    for cf_range in ws.conditional_formatting._cf_rules:
        for rule in ws.conditional_formatting._cf_rules[cf_range]:
            formulas.extend(rule.formula)
    assert any("F2>E2" in formula for formula in formulas)


def test_phase_column_is_merged(tmp_path: Path) -> None:
    path = build_workbook(tmp_path / "plan.xlsx")
    ws = load_workbook(path).active
    merged = {str(range_) for range_ in ws.merged_cells.ranges}
    assert "A2:A30" in merged
    assert "A31:A45" in merged
    assert "B2:B6" in merged
    assert "B7:B12" in merged
    assert ws["A2"].value == "第一期"
    assert ws["A31"].value == "第二期"
    assert ws["B2"].value == "前期准备"
    assert ws["B7"].value == "PAB060"


def test_source_covers_both_phases(tmp_path: Path) -> None:
    path = build_workbook(tmp_path / "plan.xlsx")
    ws = load_workbook(path).active
    names = [ws.cell(row, 3).value for row in range(2, 46) if ws.cell(row, 3).value]
    for model in ("PAB060", "PAB090", "PAB115", "PAB142", "PAB180", "PAB220"):
        assert any(str(name).startswith(f"{model}·") for name in names)
    assert names[0] == "一期零件图纸整理归档"
    assert names[-1] == "PAB220·双级样机"


def _cell_date(value):
    if value is None:
        return None
    if isinstance(value, date) and not hasattr(value, "hour"):
        return value
    if hasattr(value, "date"):
        return value.date()
    return None


def test_timeline_is_on_main_sheet_below_table(tmp_path: Path) -> None:
    ws = load_workbook(build_workbook(tmp_path / "plan.xlsx")).active
    below = list(ws.iter_rows(min_row=47, max_row=120, max_col=9))
    titles = [str(c.value) for row in below for c in row if c.value]
    assert any("目标时间轴" in t for t in titles)
    assert any("实际时间轴" in t for t in titles)
    assert any(c.value == "▶" for row in below for c in row)
    assert not ws._charts, "do not use the broken scatter chart"


def test_timeline_segments_are_chronological_unique_dates(tmp_path: Path) -> None:
    ws = load_workbook(build_workbook(tmp_path / "plan.xlsx")).active
    dates = []
    for row in ws.iter_rows(min_row=47, max_row=90, min_col=1, max_col=9):
        row_dates = [d for d in (_cell_date(c.value) for c in row) if d]
        if len(row_dates) >= 3:
            dates.extend(row_dates)
            break
    assert dates, "timeline date row should sit below the task table"
    assert dates[0] == date(2026, 8, 24)
    assert dates == sorted(dates)
    assert len(dates) == len(set(dates))
    assert len(dates) <= 9

